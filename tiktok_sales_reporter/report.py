"""
Excel 报表生成模块
将各店铺销售数据汇总为格式化的 Excel 文件，并把营业额/回款整页截图嵌入各店铺子表
"""

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 图片嵌入需要 Pillow，缺失时优雅降级
try:
    from openpyxl.drawing.image import Image as XLImage
    _IMAGE_OK = True
except Exception:
    _IMAGE_OK = False

CONFIG_PATH = Path(__file__).parent / "config" / "shops.json"

# 国旗 emoji 映射
FLAG = {"MY": "🇲🇾", "TH": "🇹🇭", "VN": "🇻🇳", "PH": "🇵🇭"}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
ALT_FILL = PatternFill("solid", fgColor="D6E4F0")
SUCCESS_FILL = PatternFill("solid", fgColor="E8F5E9")
FAIL_FILL = PatternFill("solid", fgColor="FFEBEE")

THIN_BORDER = Border(
    left=Side(style="thin", color="BDBDBD"),
    right=Side(style="thin", color="BDBDBD"),
    top=Side(style="thin", color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)

# 嵌入截图时缩放到的目标宽度（像素）
IMG_TARGET_WIDTH = 1100


def _cell(ws, row, col, value, bold=False, fill=None, align="center",
          font_color="000000", size=11, wrap=False, number_format=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, color=font_color, size=size)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = THIN_BORDER
    if fill:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format
    return cell


def _safe_sheet_name(name: str, used: set) -> str:
    """生成合法且唯一的工作表名（<=31 字符，去掉非法字符）"""
    clean = re.sub(r'[\[\]:*?/\\]', '-', name)[:28]
    candidate = clean or "shop"
    i = 1
    while candidate in used:
        candidate = f"{clean[:25]}_{i}"
        i += 1
    used.add(candidate)
    return candidate


def _embed_screenshot(ws, path: str, anchor_row: int) -> int:
    """把截图嵌入工作表，返回图片占用后下一个可用行号"""
    if not (_IMAGE_OK and path and Path(path).exists()):
        _cell(ws, anchor_row, 1, "（截图不可用，请确认已安装 Pillow: pip install pillow）",
              align="left", font_color="C00000")
        return anchor_row + 2
    try:
        img = XLImage(path)
        if img.width and img.width > IMG_TARGET_WIDTH:
            ratio = IMG_TARGET_WIDTH / float(img.width)
            img.width = IMG_TARGET_WIDTH
            img.height = int(img.height * ratio)
        ws.add_image(img, f"A{anchor_row}")
        # 估算图片占多少行（每行约 18px）
        rows_used = int(img.height / 18) + 2
        return anchor_row + rows_used
    except Exception as e:
        _cell(ws, anchor_row, 1, f"（截图嵌入失败: {e}）", align="left", font_color="C00000")
        return anchor_row + 2


def generate_report(results: list, output_dir: str = None) -> str:
    """生成 Excel 报表，返回文件路径。results 为 scraper.run_all_shops() 的返回值"""
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8-sig"))
    if output_dir is None:
        output_dir = Path(__file__).parent / config["settings"].get("reports_dir", "reports")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    today = datetime.now().strftime("%Y-%m-%d")
    filename = output_dir / f"TikTok销售日报_{today}.xlsx"
    if filename.exists():
        filename = output_dir / f"TikTok销售日报_{today}_{datetime.now().strftime('%H%M%S')}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"

    # ── 列宽（宽松，避免挤压）──
    col_widths = [6, 12, 30, 20, 18, 20, 22]
    headers = ["序号", "国家", "店铺名称", "近30天 GMV", "近30天售出件数", "可提现金额", "状态/详情"]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ── 标题行 ──
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"TikTok 多店铺销售报表（近30天）  |  {today}"
    title_cell.font = Font(bold=True, color="FFFFFF", size=16)
    title_cell.fill = HEADER_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # ── 生成时间 ──
    ws.merge_cells("A2:G2")
    ts_cell = ws["A2"]
    ts_cell.value = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  共 {len(results)} 个店铺"
    ts_cell.font = Font(color="FFFFFF", size=10)
    ts_cell.fill = SUBHEADER_FILL
    ts_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # ── 表头 ──
    for col, h in enumerate(headers, 1):
        _cell(ws, 3, col, h, bold=True, fill=SUBHEADER_FILL, font_color="FFFFFF", size=11, wrap=True)
    ws.row_dimensions[3].height = 40

    # ── 数据行 + 每个店铺的截图子表 ──
    success_count = 0
    used_names = {"汇总", "使用说明"}

    for idx, r in enumerate(results, 1):
        row = idx + 3
        fill = SUCCESS_FILL if r["status"] == "success" else FAIL_FILL
        if idx % 2 == 0 and r["status"] == "success":
            fill = ALT_FILL

        country = r["country"]
        flag = FLAG.get(country, "")

        _cell(ws, row, 1, idx, fill=fill)
        _cell(ws, row, 2, f"{flag} {country}", fill=fill)
        _cell(ws, row, 3, r["shop_name"], fill=fill, align="left")
        _cell(ws, row, 4, r.get("today_gmv", "N/A"), fill=fill, align="right")
        _cell(ws, row, 5, r.get("today_items_sold", "N/A"), fill=fill)
        _cell(ws, row, 6, r.get("available_to_withdraw", "N/A"), fill=fill, align="right")

        # 状态 / 截图详情链接
        has_shot = r.get("analytics_screenshot") or r.get("finance_screenshot")
        if r["status"] == "success" and has_shot:
            sheet_name = _safe_sheet_name(f"{country}-{r['shop_name']}", used_names)
            sub = wb.create_sheet(sheet_name)
            _build_shot_sheet(sub, r)

            link_cell = _cell(ws, row, 7, "✅ 成功（点此看截图）", fill=fill, font_color="0563C1")
            link_cell.hyperlink = f"#'{sheet_name}'!A1"
            link_cell.font = Font(color="0563C1", underline="single", size=11)
        else:
            status_text = "✅ 成功" if r["status"] == "success" else f"❌ {r.get('error', '失败')[:24]}"
            _cell(ws, row, 7, status_text, fill=fill)

        ws.row_dimensions[row].height = 24

        if r["status"] == "success":
            success_count += 1

    # ── 汇总区域（标签跨 A:B，数值跨 C:G，避免窄列挤压）──
    s = len(results) + 5
    ws.merge_cells(f"A{s}:G{s}")
    _cell(ws, s, 1, "汇总统计", bold=True, fill=HEADER_FILL, font_color="FFFFFF", size=12)
    ws.row_dimensions[s].height = 30

    stats = [
        ("总店铺数", str(len(results))),
        ("成功抓取", f"{success_count} 个"),
        ("失败店铺", f"{len(results) - success_count} 个"),
        ("覆盖国家", ", ".join(sorted(set(r["country"] for r in results if r["status"] == "success"))) or "无"),
    ]
    for i, (label, value) in enumerate(stats):
        rr = s + 1 + i
        fill = ALT_FILL if i % 2 == 0 else None
        ws.merge_cells(f"A{rr}:B{rr}")
        _cell(ws, rr, 1, label, bold=True, fill=fill, align="left")
        ws.merge_cells(f"C{rr}:G{rr}")
        _cell(ws, rr, 3, value, fill=fill, align="left")
        ws.row_dimensions[rr].height = 22

    ws.freeze_panes = "A4"

    # ── 使用说明 sheet ──
    ws_note = wb.create_sheet("使用说明")
    notes = [
        "TikTok 多店铺销售报表 - 使用说明",
        "",
        "1. GMV = Gross Merchandise Value（商品交易总额），数据为近30天",
        "2. 点击「汇总」表里每个店铺的“点此看截图”，可跳到该店铺的整页截图",
        "3. 每个店铺子表包含：营业额页面（含30天曲线）+ 回款页面 的整页截图",
        "4. 状态「需要重新登录」= 该 Chrome 窗口的 TikTok 登录已过期，请在窗口里重新登录",
        "",
        f"报表生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    for i, note in enumerate(notes, 1):
        ws_note.cell(row=i, column=1, value=note)
    ws_note.column_dimensions["A"].width = 80

    wb.save(str(filename))
    return str(filename)


def _build_shot_sheet(ws, r: dict):
    """在子表里放营业额 + 回款两张整页截图"""
    ws.column_dimensions["A"].width = 30
    ws.merge_cells("A1:H1")
    _cell(ws, 1, 1, f"{r['shop_name']} ({r['country']})  —  近30天数据截图",
          bold=True, fill=HEADER_FILL, font_color="FFFFFF", size=13, align="left")
    ws.row_dimensions[1].height = 30

    row = 3
    _cell(ws, row, 1, f"① 营业额（GMV={r.get('today_gmv','N/A')}  售出={r.get('today_items_sold','N/A')}）",
          bold=True, fill=SUBHEADER_FILL, font_color="FFFFFF", align="left")
    ws.merge_cells(f"A{row}:H{row}")
    row += 1
    row = _embed_screenshot(ws, r.get("analytics_screenshot", ""), row)

    row += 2
    _cell(ws, row, 1, f"② 回款（可提现={r.get('available_to_withdraw','N/A')}）",
          bold=True, fill=SUBHEADER_FILL, font_color="FFFFFF", align="left")
    ws.merge_cells(f"A{row}:H{row}")
    row += 1
    _embed_screenshot(ws, r.get("finance_screenshot", ""), row)
